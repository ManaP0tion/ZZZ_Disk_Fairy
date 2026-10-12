"""ZZZ 세팅 시트(xlsx) -> extension/data/benchmark.js 변환기.

사용: python3 tools/convert.py data/sheet.xlsx extension/data/benchmark.js
"""
import json, re, sys
import openpyxl

TIERS = {
    # 탭 이름: (등급 키, 표시명, 컬럼 매핑(1-based))
    "응애": ("basic", "응애", dict(spec=3, wengine=5, set4=6, set2=7, m4=8, m5=9, m6=10, subs=11,
                                  targets=[12, 13], crit=14, note=15)),
    "세팅": ("good", "이잘키", dict(spec=3, wengine=6, set4=7, set2=8, m4=9, m5=10, m6=11, subs=12,
                                  targets=[14], crit=15, note=16)),
    "악귀": ("max", "악귀", dict(set4=2, set2=3, targets=[4], crit=5, note=6)),
}

SET_FIX = {"쇼크스타 디스크": "쇼크스타 디스코"}
# W-엔진 표기를 위키/게임 이름에 맞춘다 (시트 오타, 로마 숫자)
ENGINE_FIX = {"복을 뽐는 맹호": "복을 뿜는 맹호", "II형": "Ⅱ형", "VI형": "Ⅵ형"}


def parse_engines(s):
    """엔진 목록과 메모 문장("그 외에 ... 사용 가능", "(딜세팅 용)")을 나눈다."""
    engines, notes = [], []
    for t in split_list(s):
        if t.startswith("(") or re.search(r"가능|그 외| or ", t):
            notes.append(t)
            continue
        for a, b in ENGINE_FIX.items():
            t = t.replace(a, b)
        engines.append(t)
    return engines, " ".join(notes)

TARGET_KEYS = [  # (시트 표기 정규식, 키)
    (r"공(?:격력)?", "ATK"), (r"체력|HP", "HP"), (r"이상마(?:스터리)?", "AP"),
    (r"관입력", "SHEER"), (r"충격력|충격", "IMP"), (r"방(?:어력)?", "DEF"),
    (r"이상 ?장악력", "AM"), (r"에너지(?: 회복)?", "ER"),
]


def cell(ws, r, c):
    v = ws.cell(r, c).value if c else None
    return "" if v is None else str(v).strip()


def split_list(s):
    out = []
    for t in re.split(r"[\n/,]", s):
        t = t.strip()
        if t:
            out.append(t)
    return out


def parse_subs(s):
    subs = []
    for t in split_list(s):
        low = "후순위" in t
        name = re.sub(r"\(.*?\)", "", t).strip()
        if not name:
            continue
        subs.append({"name": name, "weight": 0.5 if low else 1.0, "raw": t})
    return subs


MAIN_ABBR = {"치": "치명", "관": "관통률", "공": "공격력", "속": "속성피해", "이": "이상 마스터리", "에": "에너지 회복", "충": "충격력", "체": "체력", "방": "방어력"}


def parse_cond(text):
    """괄호 조건 해석. 평가 가능한 형태면 type을 붙이고, 아니면 note로 남긴다."""
    t = text.strip().strip("()").strip()
    m = re.match(r"^(\d)번(?: ?주옵)? ?(.+)$", t)
    if m:
        return {"type": "main", "slot": int(m.group(1)), "main": m.group(2).strip(), "text": t}
    m = re.match(r"^456 ?([가-힣]{3})$", t)
    if m:
        return {"type": "mains456", "mains": [MAIN_ABBR.get(ch, ch) for ch in m.group(1)], "text": t}
    m = re.match(r"^(\d)돌(?:파)? ?(이상|이하)?", t)
    if m and not t.startswith("w/"):
        return {"type": "rank", "n": int(m.group(1)), "op": "<=" if m.group(2) == "이하" else ">=", "text": t}
    m = re.match(r"^(\d)?\s*([가-힣/]+?)\s*(\d)?셋?(?: ?착용)?$", t)
    if m and (m.group(1) or m.group(3)):
        pc = int(m.group(1) or m.group(3))
        return {"type": "set", "pc": pc, "sets": [x for x in m.group(2).split("/") if x], "text": t}
    return {"type": "note", "text": t}


def parse_targets(texts):
    """'공 3950~4000', '이상마 470+', '(5번 주옵 공퍼)' 같은 줄을 구조화."""
    res = []
    for text in texts:
        for line in [l.strip() for l in text.split("\n") if l.strip()]:
            if re.match(r"^\(.*\)$", line) and res:
                res[-1]["cond"] = parse_cond(line)
                continue
            for pat, key in TARGET_KEYS:
                m = re.match(rf"^(?:{pat})\s*([\d,]+)\s*(?:~\s*([\d,]+))?\s*(\+)?", line)
                if m:
                    res.append({"stat": key, "min": int(m.group(1).replace(",", "")),
                                "max": int(m.group(2).replace(",", "")) if m.group(2) else None,
                                "raw": line})
                    break
    return res


def parse_crit(text):
    res = []
    for line in [l.strip() for l in text.split("\n") if l.strip()]:
        if line.startswith("(") and res:
            # "(1돌 이상은" 처럼 괄호가 다음 줄로 이어지는 경우는 메모로만 남긴다
            if line.endswith(")"):
                res[-1]["cond"] = parse_cond(line)
            else:
                res[-1]["note"] = line
                skip_close = True
            continue
        if line.endswith(")") and "(" not in line:
            continue  # 괄호 이어지는 줄 ("80 이상)", "치확 38)")
        m = re.match(r"^(\d+(?:\.\d+)?)\s*/\s*(\d+(?:\.\d+)?)", line)
        if m:
            res.append({"CR": float(m.group(1)), "CD": float(m.group(2)), "raw": line})
            continue
        m = re.search(r"(\d+(?:\.\d+)?)", line)
        if m and ("치확" in text or "기준" in line):
            res.append({"CR": float(m.group(1)), "CD": None, "raw": line})
    return res


def find_name(ws, r):
    for rr in range(r, min(r + 5, ws.max_row + 1)):
        v = cell(ws, rr, 1)
        if v:
            return v
    return None


def parse_tab(ws, cols):
    key_col = cols.get("spec") or cols["set4"]
    rows = []
    for r in range(6, ws.max_row + 1):
        if not cell(ws, r, key_col):
            continue
        name = find_name(ws, r)
        if not name or name in ("캐릭명",):
            continue
        e = {"name": name}
        if "spec" in cols:
            e["spec"] = cell(ws, r, cols["spec"])
        if "wengine" in cols:
            e["wengine"], e["wengineNote"] = parse_engines(cell(ws, r, cols["wengine"]))
        # 괄호 메모("(w/ 다이아린)")를 먼저 지운 뒤 나눈다. 시트 오타는 SET_FIX로 교정.
        strip = lambda v: re.sub(r"\([^)]*\)", "", v, flags=re.S)
        e["set4"] = [SET_FIX.get(t.strip(), t.strip()) for t in split_list(strip(cell(ws, r, cols["set4"])))]
        e["set2"] = [SET_FIX.get(t.strip(), t.strip()) for t in split_list(strip(cell(ws, r, cols["set2"])))]
        e["set4"] = [t for t in e["set4"] if t]
        e["set2"] = [t for t in e["set2"] if t]
        if "m4" in cols:
            e["mains"] = {s: [t for t in (re.sub(r"\(.*?\)", "", x).strip() for x in cell(ws, r, cols[f"m{s}"]).split("\n")) if t]
                          for s in (4, 5, 6)}
            e["subs"] = parse_subs(cell(ws, r, cols["subs"]))
        e["targets"] = parse_targets([cell(ws, r, c) for c in cols["targets"]])
        e["targetsRaw"] = " / ".join(cell(ws, r, c).replace("\n", " ") for c in cols["targets"] if cell(ws, r, c))
        e["crit"] = parse_crit(cell(ws, r, cols["crit"]))
        e["critRaw"] = cell(ws, r, cols["crit"]).replace("\n", " ")
        e["note"] = cell(ws, r, cols["note"])
        rows.append(e)
    return rows


def main(src, dst):
    wb = openpyxl.load_workbook(src)
    out = {"version": wb["세팅"]["B1"].value, "tiers": {}}
    for tab, (key, label, cols) in TIERS.items():
        rows = parse_tab(wb[tab], cols)
        out["tiers"][key] = {"label": label, "sheet": tab, "title": cell(wb[tab], 1, 2) or cell(wb[tab], 1, 1), "agents": rows}
        print(f"{tab}: {len(rows)}명")
    js = "// 자동 생성: tools/convert.py\nwindow.ZZZ_BENCHMARK = " + json.dumps(out, ensure_ascii=False, indent=1) + ";\n"
    open(dst, "w", encoding="utf-8").write(js)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
